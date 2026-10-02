"""Section requirements come from the Reference Profile for the selected DOCX."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from format.nonconformance import find_nonconformances, write_nonconformance_file
from reference_profile.sections import canonical_heading, contains_heading, required_section_headings

HISTORY = ROOT / "tests" / "fixtures" / "synthetic_reference.docx"
ENGLISH = ROOT / "input" / "Reference Exam English 10 WORD.docx"


def _teacher(path: Path, sections: list[str]) -> None:
    doc = Document()
    doc.add_paragraph("Northwind Academy")
    doc.add_paragraph("Class: 7")
    doc.add_paragraph("Subject: History")
    for heading in sections:
        doc.add_paragraph(heading)
    doc.save(str(path))


def test_english_hyphenated_heading_matches_profile_heading():
    headings = required_section_headings(ENGLISH)
    assert canonical_heading("Section - A", headings) == "Section A"
    assert canonical_heading("Section-C", headings) == "Section C"
    assert canonical_heading("Part II", headings) is None
    assert contains_heading("Part II and Part III", "Part I") is False
    assert contains_heading("Part I\nPart II", "Part I") is True


def test_history_sections_present_have_no_section_errors(tmp_path):
    teacher = tmp_path / "history.docx"
    _teacher(teacher, ["Part I", "Part II", "Part III"])
    assert find_nonconformances(teacher, HISTORY) == []


def test_history_missing_part_ii_names_only_that_section(tmp_path):
    teacher = tmp_path / "history-missing.docx"
    _teacher(teacher, ["Part I", "Part III"])
    errors = find_nonconformances(teacher, HISTORY)
    assert [item.error for item in errors] == ["Part II is missing."]
    assert errors[0].fix == "Add Part II in the correct position, following the Reference Exam."
    assert errors[0].place == "before-section:Part III"
    assert all("Section A" not in item.error and "Section B" not in item.error for item in errors)
    assert all("Section C" not in item.error and "Section D" not in item.error for item in errors)

    output = tmp_path / "corrections.docx"
    write_nonconformance_file(teacher, output, errors, HISTORY)
    texts = [paragraph.text.strip() for paragraph in Document(str(output)).paragraphs if paragraph.text.strip()]
    error_at = next(index for index, text in enumerate(texts) if text.startswith("ERROR:"))
    part_iii = next(index for index, text in enumerate(texts) if text == "Part III")
    assert error_at < part_iii
    assert "Part II is missing." in texts[error_at]
    assert not any(text == "Section A" or text == "Section B" or text == "Section C" or text == "Section D" for text in texts)
