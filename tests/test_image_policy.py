"""Image requirements come from the Reference Profile for the selected DOCX."""

from __future__ import annotations

import struct
import sys
import zipfile
import zlib
from io import BytesIO
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core.integrity import integrity
from core.models import Block
from core.number import assign_labels
from core.validate import validate
from format.format_integrity import check_format
from format.word_formatter import format_exam
from reference_profile.images import image_policy_from_profile, load_image_policy
from reference_profile.numbering import load_numbering_policy
from reference_profile.preprocess import build_profile, write_profile

ENGLISH = ROOT / "input" / "Reference Exam English 10 WORD.docx"
HISTORY = ROOT / "tests" / "fixtures" / "synthetic_reference.docx"


def _png(color: bytes) -> bytes:
    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    raw = b"\x00" + color
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw))
        + chunk(b"IEND", b"")
    )


def _block(source_id, order, text, kind, parent=None, media=None):
    return Block(
        source_id=source_id,
        source_order=order,
        original_text=text,
        block_type=kind,
        parent_id=parent,
        media_part=media,
        host_paragraph_id=parent if media else None,
        confidence=1.0,
        reason="test",
    )


def _history_blocks():
    return [
        _block("p0001", 1, "Class: 7", "metadata"),
        _block("p0002", 2, "Subject: History", "metadata"),
        _block("p0003", 3, "Part I", "section"),
        _block("p0004", 4, "1. Name one river-valley civilization.", "major_question"),
        _block("p0005", 5, "Part II", "section"),
        _block("p0006", 6, "2. Why did cities need walls?", "major_question"),
        _block("p0007", 7, "Part III", "section"),
        _block("p0008", 8, "3. Describe a marketplace.", "major_question"),
    ]


def _with_images(blocks, count):
    copied = [
        Block(
            source_id=block.source_id,
            source_order=block.source_order,
            original_text=block.original_text,
            block_type=block.block_type,
            parent_id=block.parent_id,
            media_part=block.media_part,
            host_paragraph_id=block.host_paragraph_id,
            confidence=block.confidence,
            reason=block.reason,
        )
        for block in blocks
    ]
    for index in range(1, count + 1):
        copied.append(
            _block(
                f"img{index:04d}",
                100 + index,
                "",
                "image",
                parent="p0004",
                media=f"media/image{index}.png",
            )
        )
    return copied


def _media(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return [name for name in archive.namelist() if name.startswith("word/media/")]


def test_expected_image_count_comes_from_the_profile():
    english = load_image_policy(ENGLISH)
    history = load_image_policy(HISTORY)
    assert english.count == 1
    assert len(english.items) == 1
    assert history.count == 0
    assert history.items == ()

    numbering = load_numbering_policy(HISTORY)
    empty = _history_blocks()
    assign_labels(empty, numbering)
    assert validate(empty, True, numbering, history) == []
    presence = next(item for item in integrity(empty, empty, history) if item[0] == "image_presence")
    assert presence[1] is True

    english_numbering = load_numbering_policy(ENGLISH)
    assign_labels(empty, english_numbering)
    problems = validate(empty, True, english_numbering, english)
    assert problems == ["Expected 1 image block(s), found 0."]
    missing = next(item for item in integrity(empty, empty, english) if item[0] == "image_presence")
    assert missing[1] is False

    one = _with_images(_history_blocks(), 1)
    assign_labels(one, english_numbering)
    assert validate(one, True, english_numbering, english) == []
    kept = next(item for item in integrity(one, one, english) if item[0] == "image_presence")
    assert kept[1] is True


def test_two_image_profile_is_not_forced_to_one(tmp_path):
    path = tmp_path / "two-images.docx"
    doc = Document()
    doc.add_paragraph("Part I")
    doc.add_paragraph("1. First prompt.")
    doc.add_picture(BytesIO(_png(b"\xff\x00\x00")))
    doc.add_paragraph("2. Second prompt.")
    doc.add_picture(BytesIO(_png(b"\x00\x00\xff")))
    doc.save(str(path))

    profile = build_profile(path)
    assert profile["images"]["status"] == "observed"
    assert profile["images"]["count"] == 2
    assert len(profile["images"]["items"]) == 2
    write_profile(path, tmp_path / "profiles")
    policy = load_image_policy(path, tmp_path / "profiles")
    assert policy.count == 2
    assert image_policy_from_profile(profile).count == 2

    numbering = load_numbering_policy(HISTORY)
    two = _with_images(_history_blocks(), 2)
    assign_labels(two, numbering)
    assert validate(two, True, numbering, policy) == []
    one = _with_images(_history_blocks(), 1)
    assign_labels(one, numbering)
    assert validate(one, True, numbering, policy) == ["Expected 2 image block(s), found 1."]


def test_history_zero_images_survive_a_second_format(tmp_path):
    numbering = load_numbering_policy(HISTORY)
    image_policy = load_image_policy(HISTORY)
    assert image_policy.count == 0
    blocks = _history_blocks()
    assign_labels(blocks, numbering)
    once = tmp_path / "once.docx"
    format_exam(blocks, HISTORY, once, HISTORY)
    assert _media(once) == []
    first = check_format(blocks, HISTORY, once, HISTORY)
    assert not any("image" in problem.lower() for problem in first)

    twice = tmp_path / "twice.docx"
    format_exam(blocks, once, twice, HISTORY)
    assert _media(twice) == []
    second = check_format(blocks, once, twice, HISTORY)
    assert not any("image" in problem.lower() for problem in second)
