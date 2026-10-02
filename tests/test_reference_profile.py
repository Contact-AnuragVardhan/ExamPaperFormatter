"""Reference profiles are measured from DOCX files."""

from __future__ import annotations

import ast
import hashlib
import json
import re
import sys
import zipfile
from collections import Counter
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from reference_profile import PROFILE_VERSION, build_profile, write_profile

ENGLISH = ROOT / "input" / "Reference Exam English 10 WORD.docx"
SYNTHETIC = ROOT / "tests" / "fixtures" / "synthetic_reference.docx"
TEACHER = ROOT / "input" / "Teacher Exam English 10 Actual.docx"
WP_EXTENT = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent"
A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"


def _english():
    return build_profile(ENGLISH)


def _synthetic():
    return build_profile(SYNTHETIC)


def test_profile_version_and_hash_and_source_is_unchanged():
    before = ENGLISH.read_bytes()
    profile = _english()
    assert ENGLISH.read_bytes() == before
    assert profile["profile_version"] == PROFILE_VERSION == 1
    assert profile["source_reference"] == ENGLISH.name
    assert profile["source_hash"] == {
        "algorithm": "sha256",
        "hex": hashlib.sha256(before).hexdigest(),
    }
    assert build_profile(ENGLISH) == profile


def test_english_profile_matches_the_reference_docx():
    profile = _english()
    doc = Document(str(ENGLISH))
    section = doc.sections[0]
    page = profile["appearance"]["page"]
    margins = profile["appearance"]["margins"]
    assert page["width_twip"] == int(round(section.page_width.twips))
    assert page["height_twip"] == int(round(section.page_height.twips))
    assert margins["left_twip"] == int(round(section.left_margin.twips))
    assert margins["right_twip"] == int(round(section.right_margin.twips))
    assert margins["top_twip"] == int(round(section.top_margin.twips))
    assert margins["bottom_twip"] == int(round(section.bottom_margin.twips))

    fonts = Counter()
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.text and run.font.name:
                fonts[run.font.name] += 1
    assert profile["appearance"]["predominant_font"]["value"] == fonts.most_common(1)[0][0]
    assert profile["appearance"]["predominant_font"]["status"] == "observed"

    headings = [item["heading"] for item in profile["structure"]["sections"]["items"]]
    assert headings == ["Section A", "Section B", "Section C", "Section D"]
    for item in profile["structure"]["sections"]["items"]:
        assert doc.paragraphs[item["paragraph_index"] - 1].text.strip() == item["heading"]
    section_paragraph = doc.paragraphs[profile["structure"]["sections"]["items"][0]["paragraph_index"] - 1]
    assert profile["appearance"]["roles"]["section_heading"]["font_size_pt"]["value"] == int(round(section_paragraph.runs[0].font.size.pt))
    assert profile["appearance"]["roles"]["section_heading"]["bold"]["value"] is True

    numbers = []
    for paragraph in doc.paragraphs:
        match = re.match(r"Q\s*(\d+)", paragraph.text.strip())
        if match:
            numbers.append(int(match.group(1)))
    assert profile["structure"]["major_questions"]["status"] == "observed"
    assert profile["structure"]["major_questions"]["count"] == len(set(numbers))
    assert profile["structure"]["major_questions"]["numbers"] == list(range(1, len(set(numbers)) + 1))
    first_major = next(paragraph.text.strip() for paragraph in doc.paragraphs if re.match(r"Q\s*\d+", paragraph.text.strip()))
    assert first_major.startswith(profile["numbering"]["major"]["examples"][0])
    for example in profile["numbering"]["major"]["examples"]:
        assert any(paragraph.text.strip().startswith(example) for paragraph in doc.paragraphs)

    letters = []
    for paragraph in doc.paragraphs:
        match = re.match(r"Q\s*9\s*\(\s*([A-Za-z])\s*\)", paragraph.text.strip())
        if match:
            letters.append(match.group(1))
    assert letters
    assert profile["numbering"]["branch"]["letters"] == letters
    for example in profile["numbering"]["subquestion"]["examples"]:
        assert any(paragraph.text.strip().startswith(example) for paragraph in doc.paragraphs)
    choice_examples = [
        example
        for kind in profile["numbering"]["choice"]["kinds"]
        for example in kind["examples"]
    ]
    assert any(example.startswith("A)") for example in choice_examples)
    assert any(paragraph.text.strip().startswith("A)") for paragraph in doc.paragraphs)

    first_text = " ".join(next(paragraph.text for paragraph in doc.paragraphs if paragraph.text.strip()).split())
    assert profile["header"]["lines"][0]["text"] == first_text
    assert profile["identity"]["subject"]["evidence"] == "Subject: English"
    assert profile["identity"]["class_grade"]["evidence"] == "Grade: 10"
    assert profile["identity"]["subject"]["value"] == "English"
    assert profile["identity"]["class_grade"]["value"] == "10"
    assert profile["structure"]["syllabus"]["status"] == "observed"
    assert doc.paragraphs[profile["structure"]["syllabus"]["paragraph_index"] - 1].text.strip() == "SYLLABUS"
    assert profile["structure"]["syllabus"]["location"] == "before_first_section"
    assert profile["numbering"]["alternative"]["status"] == "not_determined"

    image_paragraphs = []
    for index, paragraph in enumerate(doc.paragraphs, start=1):
        blips = list(paragraph._p.iter(A_BLIP))
        if not blips:
            continue
        extent = next(paragraph._p.iter(WP_EXTENT), None)
        image_paragraphs.append((index, extent))
    assert profile["images"]["count"] == len(image_paragraphs) == 1
    item = profile["images"]["items"][0]
    assert item["paragraph_index"] == image_paragraphs[0][0]
    assert item["width_emu"] == int(image_paragraphs[0][1].get("cx"))
    assert item["height_emu"] == int(image_paragraphs[0][1].get("cy"))
    previous = next(
        paragraph.text.strip()
        for paragraph in reversed(doc.paragraphs[: item["paragraph_index"] - 1])
        if paragraph.text.strip()
    )
    assert previous.startswith(item["host_marker"])

    with zipfile.ZipFile(ENGLISH) as archive:
        styles = etree.fromstring(archive.read("word/styles.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    spacing = styles.find("w:docDefaults/w:pPrDefault/w:pPr/w:spacing", ns)
    fonts = styles.find("w:docDefaults/w:rPrDefault/w:rPr/w:rFonts", ns)
    assert profile["appearance"]["document_defaults"]["line_spacing"]["value"] == int(spacing.get(qn("w:line")))
    if fonts.get(qn("w:ascii")):
        assert profile["appearance"]["document_defaults"]["font_name"]["value"] == fonts.get(qn("w:ascii"))
    else:
        assert profile["appearance"]["document_defaults"]["font_name"]["status"] == "not_determined"

    blob = json.dumps(profile)
    assert "Lencho" not in blob
    assert "First Flight" not in blob


def test_synthetic_profile_is_not_the_english_profile():
    english = _english()
    profile = _synthetic()
    doc = Document(str(SYNTHETIC))
    headings = [item["heading"] for item in profile["structure"]["sections"]["items"]]
    assert headings == ["Part I", "Part II", "Part III"]
    assert "Section A" not in headings
    for item in profile["structure"]["sections"]["items"]:
        assert doc.paragraphs[item["paragraph_index"] - 1].text.strip() == item["heading"]

    numbers = []
    for paragraph in doc.paragraphs:
        match = re.match(r"(\d+)\s*[\.)]", paragraph.text.strip())
        if match:
            numbers.append(int(match.group(1)))
    assert profile["structure"]["major_questions"]["count"] == len(set(numbers)) == 4
    assert all("Q" not in pattern for pattern in profile["numbering"]["major"]["patterns"])
    assert profile["numbering"]["subquestion"]["examples"] == ["a.", "b."]
    assert profile["numbering"]["choice"]["status"] == "not_determined"
    assert profile["numbering"]["branch"]["status"] == "not_determined"
    assert profile["numbering"]["alternative"]["status"] == "not_determined"
    assert profile["structure"]["syllabus"]["status"] == "not_determined"
    assert profile["images"]["count"] == 0

    fonts = Counter(run.font.name for paragraph in doc.paragraphs for run in paragraph.runs if run.text and run.font.name)
    assert profile["appearance"]["predominant_font"]["value"] == fonts.most_common(1)[0][0]
    assert profile["appearance"]["predominant_font"]["value"] != english["appearance"]["predominant_font"]["value"]
    assert profile["appearance"]["page"]["width_twip"] == int(round(doc.sections[0].page_width.twips))
    assert profile["appearance"]["margins"]["left_twip"] == int(round(doc.sections[0].left_margin.twips))
    assert profile["appearance"]["page"]["width_twip"] != english["appearance"]["page"]["width_twip"]
    assert profile["appearance"]["margins"]["left_twip"] != english["appearance"]["margins"]["left_twip"]
    assert profile["structure"]["major_questions"]["count"] != english["structure"]["major_questions"]["count"]
    assert profile["images"]["count"] != english["images"]["count"]
    assert profile["identity"]["subject"]["value"] == "History"
    assert profile["identity"]["class_grade"]["value"] == "7"
    assert profile["identity"]["subject"]["value"] != english["identity"]["subject"]["value"]


def test_preprocessor_has_no_english_10_constants_and_no_runtime_imports():
    source = "\n".join(path.read_text(encoding="utf-8") for path in (ROOT / "reference_profile").glob("*.py"))
    for banned in ("Calibri", "Jalta", "Times New Roman", "Lencho", "First Flight", "Q1)", "ABCD", "0.75", "8.5", "English 10"):
        assert banned not in source
    tree = ast.parse((ROOT / "reference_profile" / "preprocess.py").read_text(encoding="utf-8"))
    modules = set()
    for node in ast.walk(tree):
        if isinstance(node, ast.Import):
            modules.update(alias.name.split(".")[0] for alias in node.names)
        elif isinstance(node, ast.ImportFrom) and node.module:
            modules.add(node.module.split(".")[0])
    assert modules.isdisjoint({"core", "format", "web", "openai", "run_core"})


def test_write_profile_json_round_trip(tmp_path):
    target = write_profile(SYNTHETIC, tmp_path)
    assert target.parent == tmp_path
    assert target.suffix == ".json"
    loaded = json.loads(target.read_text(encoding="utf-8"))
    assert loaded == build_profile(SYNTHETIC)
    assert loaded["source_hash"]["hex"][:12] in target.name


def test_runtime_section_policy_requires_a_matching_profile(tmp_path):
    from docx import Document

    from format.nonconformance import find_nonconformances
    from reference_profile.sections import ReferenceProfileError, required_section_headings
    from web.app import create_app

    assert required_section_headings(ENGLISH) == ["Section A", "Section B", "Section C", "Section D"]
    unrelated = tmp_path / "unprofiled.docx"
    Document().save(str(unrelated))
    try:
        required_section_headings(unrelated)
    except ReferenceProfileError as exc:
        assert "Refusing to assume a section sequence" in str(exc)
    else:
        raise AssertionError("A reference without a profile was accepted.")
    assert find_nonconformances(TEACHER, ENGLISH) == []
    client = create_app(tmp_path / "data").test_client()
    page = client.get("/")
    assert page.status_code == 200
    assert not (tmp_path / "data" / "reference_profiles").exists()
