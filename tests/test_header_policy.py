"""Header requirements come from the Reference Profile for the selected DOCX."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core.models import Block
from format.nonconformance import find_nonconformances
from format.word_formatter import format_exam
from reference_profile.headers import load_header_policy

ENGLISH = ROOT / "input" / "Reference Exam English 10 WORD.docx"
HISTORY = ROOT / "tests" / "fixtures" / "synthetic_reference.docx"


def _teacher(path: Path, lines: list[str]) -> None:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    doc.save(str(path))


def _block(source_id, order, text, kind):
    return Block(
        source_id=source_id,
        source_order=order,
        original_text=text,
        block_type=kind,
        confidence=1.0,
        reason="test",
    )


def test_required_header_roles_follow_each_profile():
    english = load_header_policy(ENGLISH)
    history = load_header_policy(HISTORY)
    assert [item.role for item in english] == ["class_grade", "subject", "time", "marks"]
    assert [item.label for item in english] == ["Class", "Subject", "Time", "Maximum marks"]
    assert [item.role for item in history] == ["class_grade", "subject"]
    assert english[0].matches("Class X")
    assert english[0].matches("Grade: 10")
    assert history[0].matches("Class: 7")
    assert not any(item.role == "time" or item.role == "marks" for item in history)


def test_history_header_present_has_no_header_errors(tmp_path):
    teacher = tmp_path / "history.docx"
    _teacher(
        teacher,
        ["Northwind Academy", "Term Assessment", "Class: 7", "Subject: History", "Part I", "Part II", "Part III"],
    )
    assert find_nonconformances(teacher, HISTORY) == []

    without_undetermined = tmp_path / "roles-only.docx"
    _teacher(without_undetermined, ["Class: 7", "Subject: History", "Part I", "Part II", "Part III"])
    assert find_nonconformances(without_undetermined, HISTORY) == []


def test_history_missing_subject_names_only_that_header(tmp_path):
    teacher = tmp_path / "missing-subject.docx"
    _teacher(teacher, ["Northwind Academy", "Term Assessment", "Class: 7", "Part I", "Part II", "Part III"])
    errors = find_nonconformances(teacher, HISTORY)
    assert [item.error for item in errors] == ["Subject header is missing."]
    assert errors[0].fix == "Add the Subject header in the header area, following the Reference Exam."
    joined = " ".join(item.error for item in errors)
    assert "Time header" not in joined
    assert "Maximum marks" not in joined
    assert "Class header" not in joined


def test_history_header_survives_a_second_format(tmp_path):
    lines = ["Northwind Academy", "Term Assessment", "Class: 7", "Subject: History"]
    blocks = [_block(f"p{index:04d}", index, text, "metadata") for index, text in enumerate(lines, start=1)]
    blocks.append(_block("p0005", 5, "Part I", "section"))
    once = tmp_path / "once.docx"
    format_exam(blocks, HISTORY, once, HISTORY)
    first = [paragraph.text for paragraph in Document(str(once)).paragraphs if paragraph.text.strip()]
    assert first[:4] == lines
    assert first.count("Class: 7") == 1
    assert first.count("Subject: History") == 1

    for block, text in zip(blocks, first):
        block.original_text = text
    twice = tmp_path / "twice.docx"
    format_exam(blocks, once, twice, HISTORY)
    second = [paragraph.text for paragraph in Document(str(twice)).paragraphs if paragraph.text.strip()]
    assert second == first
