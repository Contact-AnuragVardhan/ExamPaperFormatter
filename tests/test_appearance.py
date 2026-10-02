"""Page layout and role formatting come from the selected Reference Profile."""

from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Twips

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core.models import Block
from core.number import assign_labels
from format.word_formatter import format_exam
from reference_profile.appearance import load_appearance
from reference_profile.numbering import load_numbering_policy
from reference_profile.sections import load_profile_for_reference

ENGLISH = ROOT / "input" / "Reference Exam English 10 WORD.docx"
HISTORY = ROOT / "tests" / "fixtures" / "synthetic_reference.docx"


def _block(source_id, order, text, kind, parent=None):
    return Block(
        source_id=source_id,
        source_order=order,
        original_text=text,
        block_type=kind,
        parent_id=parent,
        confidence=1.0,
        reason="test",
    )


def _history_blocks():
    return [
        _block("p0001", 1, "Northwind Academy", "metadata"),
        _block("p0002", 2, "Term Assessment", "metadata"),
        _block("p0003", 3, "Class: 7", "metadata"),
        _block("p0004", 4, "Subject: History", "metadata"),
        _block("p0005", 5, "Part I", "section"),
        _block("p0006", 6, "1. Name one river-valley civilization.", "major_question"),
        _block("p0007", 7, "a. Mention where people settled.", "subquestion", parent="p0006"),
    ]


def _twips(length) -> int:
    return int(round(length.twips))


def test_profile_appearance_differs_and_document_default_font_stays_unknown():
    english = load_appearance(ENGLISH)
    history = load_appearance(HISTORY)
    english_profile = load_profile_for_reference(ENGLISH)
    history_profile = load_profile_for_reference(HISTORY)
    assert english.page_width_twip == 12240
    assert english.page_height_twip == 15840
    assert english.margin_left_twip == 1080
    assert english.predominant_font == "Calibri"
    assert english.role("section_heading").font_size_pt == 13
    assert english.role("major_question").font_size_pt == 12
    assert english_profile["appearance"]["document_defaults"]["font_name"]["status"] == "not_determined"
    assert history.page_width_twip == 11906
    assert history.page_height_twip == 16838
    assert history.margin_left_twip == 720
    assert history.predominant_font == "Times New Roman"
    assert history.role("section_heading").font_size_pt == 16
    assert history.role("major_question").font_size_pt == 14
    assert history.role("subquestion").font_size_pt == 12
    assert history_profile["appearance"]["document_defaults"]["font_name"]["status"] == "not_determined"
    assert history.page_width_twip != english.page_width_twip
    assert history.margin_left_twip != english.margin_left_twip


def test_english_output_keeps_reference_page_and_production_underlines(tmp_path):
    blocks = [
        _block("p0001", 1, "Bhagini Nivedita", "metadata"),
        _block("p0002", 2, "Quarterly Exam", "metadata"),
        _block("p0003", 3, "Class X", "metadata"),
        _block("p0004", 4, "Subject: English", "metadata"),
        _block("p0005", 5, "Time: 3 Hrs M.M.: 80", "metadata"),
        _block("p0006", 6, "Section A", "section"),
        _block("p0007", 7, "1. Read the passage.", "major_question"),
    ]
    assign_labels(blocks, load_numbering_policy(ENGLISH))
    path = tmp_path / "english.docx"
    format_exam(blocks, ENGLISH, path, ENGLISH)
    doc = Document(str(path))
    section = doc.sections[0]
    assert _twips(section.page_width) == 12240
    assert _twips(section.page_height) == 15840
    assert _twips(section.left_margin) == 1080
    heading = next(paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == "Section A")
    run = next(item for item in heading.runs if item.text)
    assert run.font.name == "Calibri"
    assert run.font.size.pt == 13
    assert run.bold is True
    assert run.underline is not None and run.underline is not False
    assert heading.paragraph_format.space_before.pt == 12
    assert heading.paragraph_format.space_after.pt == 2
    question = next(paragraph for paragraph in doc.paragraphs if paragraph.text.strip().startswith("Q1)"))
    label = next(item for item in question.runs if item.text)
    assert label.font.size.pt == 12
    assert label.bold is True
    assert label.underline is not None and label.underline is not False


def test_history_output_uses_reference_appearance_and_round_trips(tmp_path):
    blocks = _history_blocks()
    assign_labels(blocks, load_numbering_policy(HISTORY))
    once = tmp_path / "once.docx"
    format_exam(blocks, HISTORY, once, HISTORY)
    first = Document(str(once))
    section = first.sections[0]
    assert _twips(section.page_width) == 11906
    assert _twips(section.page_height) == 16838
    assert _twips(section.left_margin) == 720
    assert _twips(section.right_margin) == 720
    assert _twips(section.top_margin) == 720
    assert _twips(section.bottom_margin) == 720
    assert _twips(section.page_width) != 12240
    assert _twips(section.left_margin) != 1080

    part = next(paragraph for paragraph in first.paragraphs if paragraph.text.strip() == "Part I")
    part_run = next(run for run in part.runs if run.text)
    assert part_run.font.name == "Times New Roman"
    assert part_run.font.name != "Calibri"
    assert part_run.font.size == Twips(16 * 20)
    assert part_run.bold is True
    assert part.paragraph_format.space_before.pt == 0
    assert part.paragraph_format.space_after.pt == 6

    major = next(paragraph for paragraph in first.paragraphs if paragraph.text.strip().startswith("1."))
    major_run = next(run for run in major.runs if run.text)
    assert major_run.font.name == "Times New Roman"
    assert major_run.font.size.pt == 14

    sub = next(paragraph for paragraph in first.paragraphs if paragraph.text.strip().startswith("a."))
    sub_run = next(run for run in sub.runs if run.text)
    assert sub_run.font.name == "Times New Roman"
    assert sub_run.font.size.pt == 12

    twice = tmp_path / "twice.docx"
    format_exam(blocks, once, twice, HISTORY)
    second = Document(str(twice))
    assert _twips(second.sections[0].page_width) == _twips(section.page_width)
    assert _twips(second.sections[0].page_height) == _twips(section.page_height)
    assert _twips(second.sections[0].left_margin) == _twips(section.left_margin)
    second_part = next(paragraph for paragraph in second.paragraphs if paragraph.text.strip() == "Part I")
    second_run = next(run for run in second_part.runs if run.text)
    assert second_run.font.name == part_run.font.name
    assert second_run.font.size == part_run.font.size
    assert second_part.paragraph_format.space_before == part.paragraph_format.space_before
    assert second_part.paragraph_format.space_after == part.paragraph_format.space_after
