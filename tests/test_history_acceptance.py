"""Synthetic History 7 through the same production path as English 10."""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from format.format_integrity import check_format
from format.nonconformance import find_nonconformances, write_nonconformance_file
from format.word_formatter import format_exam
from run_core import run_pipeline
from web.app import create_app

HISTORY = ROOT / "tests" / "fixtures" / "synthetic_reference.docx"
GOOD = ROOT / "tests" / "fixtures" / "history_teacher_good.docx"
BAD = ROOT / "tests" / "fixtures" / "history_teacher_bad.docx"
ENGLISH_ONLY = (
    "Section A",
    "Section B",
    "Section C",
    "Section D",
    "SYLLABUS",
    "Jalta Sitara",
    "Maximum Marks",
    "Q1)",
)


def _texts(path: Path) -> list[str]:
    return [paragraph.text.strip() for paragraph in Document(str(path)).paragraphs if paragraph.text.strip()]


def _twips(length) -> int:
    return int(round(length.twips))


def assert_history_exam(path: Path) -> None:
    doc = Document(str(path))
    texts = [paragraph.text.strip() for paragraph in doc.paragraphs if paragraph.text.strip()]
    flat = "\n".join(texts)
    for line in (
        "Northwind Academy",
        "Term Assessment",
        "Class: 7",
        "Subject: History",
        "Part I",
        "Part II",
        "Part III",
    ):
        assert line in texts, line
    assert texts.index("Part I") < texts.index("Part II") < texts.index("Part III")
    majors = [text for text in texts if text.startswith(("1. ", "2. ", "3. ", "4. "))]
    assert [text.split()[0] for text in majors] == ["1.", "2.", "3.", "4."]
    assert any(text.startswith("a. ") for text in texts)
    assert any(text.startswith("b. ") for text in texts)
    assert not any("1. 1." in text or "a. a." in text for text in texts)
    assert not any(text.startswith("Q") for text in texts)
    for banned in ENGLISH_ONLY:
        assert banned not in flat
    assert not any(text.lower().startswith("time") for text in texts)
    assert "Calibri" not in flat
    with zipfile.ZipFile(path) as zf:
        assert not [name for name in zf.namelist() if name.startswith("word/media/")]
    section = doc.sections[0]
    assert _twips(section.page_width) == 11906
    assert _twips(section.page_height) == 16838
    assert _twips(section.left_margin) == 720
    assert _twips(section.top_margin) == 720
    part = next(paragraph for paragraph in doc.paragraphs if paragraph.text.strip() == "Part I")
    run = next(item for item in part.runs if item.text)
    assert run.font.name == "Times New Roman"
    assert run.font.size.pt == 16
    question = next(paragraph for paragraph in doc.paragraphs if paragraph.text.strip().startswith("1. "))
    question_run = next(item for item in question.runs if item.text)
    assert question_run.font.name == "Times New Roman"
    assert question_run.font.size.pt == 14


def test_history_bad_exam_returns_history_corrections(tmp_path):
    teacher = BAD
    errors = find_nonconformances(teacher, HISTORY)
    messages = [item.error for item in errors]
    assert "Subject header is missing." in messages
    assert "Part II is missing." in messages
    joined = "\n".join(item.error + "\n" + item.fix for item in errors)
    for banned in ("Time", "Maximum", "Section A", "Section B", "Section C", "Section D", "SYLLABUS", "Q1)", "image"):
        assert banned not in joined

    output = tmp_path / "corrections.docx"
    write_nonconformance_file(teacher, output, errors, HISTORY)
    doc = Document(str(output))
    notes = [paragraph for paragraph in doc.paragraphs if paragraph.text.startswith(("ERROR:", "FIX:"))]
    assert len(notes) == 4
    assert any("Subject header is missing." in paragraph.text for paragraph in notes)
    assert any("Part II is missing." in paragraph.text for paragraph in notes)
    for paragraph in notes:
        runs = [run for run in paragraph.runs if run.text]
        assert runs
        assert all(run.font.highlight_color == WD_COLOR_INDEX.YELLOW for run in runs)


def test_history_good_exam_formats(tmp_path):
    teacher = GOOD
    assert find_nonconformances(teacher, HISTORY) == []
    result = run_pipeline(teacher, HISTORY, tmp_path / "work")
    assert result["ok"], (result["problems"], result["checks"])
    output = tmp_path / "formatted.docx"
    format_exam(result["blocks"], teacher, output, HISTORY)
    problems = check_format(result["blocks"], teacher, output, HISTORY)
    assert problems == [], problems
    assert_history_exam(output)


def test_history_formatted_exam_round_trips(tmp_path):
    teacher = GOOD
    first = run_pipeline(teacher, HISTORY, tmp_path / "first")
    assert first["ok"], first["problems"]
    once = tmp_path / "once.docx"
    format_exam(first["blocks"], teacher, once, HISTORY)
    assert find_nonconformances(once, HISTORY) == []

    second = run_pipeline(once, HISTORY, tmp_path / "second")
    assert second["ok"], second["problems"]
    twice = tmp_path / "twice.docx"
    format_exam(second["blocks"], once, twice, HISTORY)
    problems = check_format(second["blocks"], once, twice, HISTORY)
    assert problems == [], problems
    assert_history_exam(twice)
    assert _texts(once) == _texts(twice)
    first_doc = Document(str(once))
    second_doc = Document(str(twice))
    assert _twips(first_doc.sections[0].page_width) == _twips(second_doc.sections[0].page_width)
    assert _twips(first_doc.sections[0].left_margin) == _twips(second_doc.sections[0].left_margin)


def test_browser_history_workflow(tmp_path):
    teacher = GOOD
    app = create_app(tmp_path / "data")
    client = app.test_client()
    uploaded = client.post(
        "/reference",
        data={
            "grade": "7",
            "subject": "History",
            "reference": (HISTORY.open("rb"), HISTORY.name),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"synthetic_reference.docx" in uploaded.data
    formatted = client.post(
        "/format",
        data={
            "grade": "7",
            "subject": "History",
            "teacher": (teacher.open("rb"), "History 7 Teacher.docx"),
        },
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert b"DOWNLOAD FORMATTED EXAM" in formatted.data
    assert b"History 7 Teacher_FORMATTED.docx" in formatted.data
    page = formatted.data.decode("utf-8")
    start = page.find("/exams/")
    exam_id = page[start + len("/exams/") : page.find("/download", start)]
    download = client.get(f"/exams/{exam_id}/download")
    assert download.status_code == 200
    output = tmp_path / "browser.docx"
    output.write_bytes(download.data)
    download.close()
    assert_history_exam(output)


class _FakeWhatsApp:
    def __init__(self):
        self.documents = []

    def send_text(self, to, text, phone_number_id=None):
        return None

    def send_document(self, to, path, filename, caption, phone_number_id=None):
        self.documents.append((to, Path(path), filename, caption, phone_number_id))


def test_whatsapp_history_formats(tmp_path, monkeypatch):
    monkeypatch.setenv("WHATSAPP_ASYNC", "false")
    teacher = GOOD
    app = create_app(tmp_path / "data")
    bot = app.config["WHATSAPP_BOT"]
    fake = _FakeWhatsApp()
    bot.client = fake
    reference = tmp_path / "data" / "references" / "active.docx"
    reference.parent.mkdir(parents=True, exist_ok=True)
    reference.write_bytes(HISTORY.read_bytes())
    session = {
        "state": "formatting",
        "incoming_dir": str(tmp_path / "incoming"),
        "original_path": str(teacher),
        "original_filename": "History 7 Teacher.docx",
        "profile_name": "Teacher",
        "phone_number_id": "12345",
        "metadata": {"class_grade": "7", "subject": "History"},
    }
    (tmp_path / "incoming").mkdir()
    bot.store.set("15551234567", session)
    bot._format_and_send("15551234567", session)
    assert len(fake.documents) == 1
    _to, _path, filename, caption, _phone = fake.documents[0]
    assert filename == "History 7 Teacher_FORMATTED.docx"
    assert "formatted exam is ready" in caption.lower()
    rows = bot.load_index()
    saved = tmp_path / "data" / "exams" / rows[0]["id"] / "formatted.docx"
    assert_history_exam(saved)
