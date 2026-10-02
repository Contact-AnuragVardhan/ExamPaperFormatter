"""Numbering labels come from the Reference Profile for the selected DOCX."""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from core.discover import _system_prompt, reference_convention_notes
from core.extract import extract_document
from core.models import Block
from core.number import assign_labels
from format.word_formatter import format_exam
from reference_profile.numbering import load_numbering_policy, observed_patterns
from reference_profile.sections import load_profile_for_reference, required_section_headings

ENGLISH = ROOT / "input" / "Reference Exam English 10 WORD.docx"
HISTORY = ROOT / "tests" / "fixtures" / "synthetic_reference.docx"


def _block(source_id, order, text, kind, parent=None, carrier=False, slots=1):
    return Block(
        source_id=source_id,
        source_order=order,
        original_text=text,
        block_type=kind,
        parent_id=parent,
        branch_carrier=carrier,
        slots=slots,
        confidence=1.0,
        reason="test",
    )


def test_english_profile_observation_differs_from_production_output():
    profile = load_profile_for_reference(ENGLISH)
    observed = observed_patterns(profile)
    assert observed["major"] == ("Q{n}.", "Q{n}({letter}).")
    assert "Q{n})" not in observed["major"]
    assert observed["subquestion"] == ("({roman})",)
    assert observed["branch"] == ("Q{n}({letter}).",)
    assert observed["choice"] == ("({letter})", "{Letter})")

    policy = load_numbering_policy(ENGLISH)
    assert policy.source == "compatibility"
    assert policy.alternative == "OR"
    assert policy.label("major", 1) == "Q1)"
    assert policy.label("branch", 1) == "A)"
    assert policy.label("subquestion", 1) == "i)"
    assert policy.label("choice", 1) == "a)"
    assert policy.strip_leading("Q1) Read") == "Read"
    assert policy.strip_leading("1.Read") == "Read"


def test_english_compatibility_numbering_matches_current_labels():
    policy = load_numbering_policy(ENGLISH)
    blocks = [
        _block("p0001", 1, "9.(a) Read the branch.", "major_question", carrier=True),
        _block("p0002", 2, "i. First point", "subquestion", parent="p0001"),
        _block("p0003", 3, "a. one b. two c. three d. four", "choice", parent="p0002", slots=4),
        _block("p0004", 4, "(B) Other branch", "branch", parent="p0001"),
        _block("p0005", 5, "i. Branch point", "subquestion", parent="p0004"),
    ]
    assign_labels(blocks, policy)
    by_id = {block.source_id: block.final_label for block in blocks}
    assert by_id["p0001"] == "Q1) A)"
    assert by_id["p0002"] == "i)"
    assert by_id["p0003"] == "a) b) c) d)"
    assert by_id["p0004"] == "B)"
    assert by_id["p0005"] == "i)"


def test_english_discovery_prompt_stays_on_the_compatibility_wording():
    policy = load_numbering_policy(ENGLISH)
    headings = required_section_headings(ENGLISH)
    prompt = _system_prompt(headings, policy)
    assert 'A leading integer may be "9.(a)"' in prompt
    assert "Its text is branch A." in prompt
    assert "i. ii. iii." in prompt
    notes = reference_convention_notes(extract_document(ENGLISH), headings, policy)
    assert notes == (
        "paragraphs=148 images=1 section_headings=4 "
        "q_prefixed_lines=17 parenthesized_roman_lines=45 "
        "uppercase_choice_lines=40 standalone_OR_lines=0"
    )


def test_history_numbering_uses_observed_markers_and_round_trips(tmp_path):
    profile = load_profile_for_reference(HISTORY)
    observed = observed_patterns(profile)
    assert observed["major"] == ("{n}.",)
    assert observed["subquestion"] == ("{letter}.",)
    assert observed["choice"] == ()
    assert observed["branch"] == ()

    policy = load_numbering_policy(HISTORY)
    assert policy.source == "observed"
    assert policy.alternative is None
    assert policy.major == "{n}."
    assert policy.subquestion == "{letter}."
    assert policy.choice is None
    assert policy.branch is None
    assert policy.label("major", 1) == "1."
    assert policy.label("subquestion", 1) == "a."
    assert policy.strip_leading("1. Name one river") == "Name one river"
    assert policy.strip_leading("a. Mention where") == "Mention where"
    assert policy.strip_leading("Q1) Name one river") == "Q1) Name one river"

    prompt = _system_prompt(["Part I", "Part II", "Part III"], policy)
    assert "1." in prompt and "a." in prompt
    assert "no branch convention" in prompt
    assert "no choice convention" in prompt
    assert "Q1)" not in prompt
    assert "branch A" not in prompt
    assert "i. ii. iii." not in prompt
    assert "OR" not in prompt
    assert "SYLLABUS" not in prompt
    assert "no syllabus heading" in prompt
    assert "no alternative convention" in prompt

    blocks = [
        _block("p0001", 1, "Northwind Academy", "metadata"),
        _block("p0002", 2, "Class: 7", "metadata"),
        _block("p0003", 3, "Subject: History", "metadata"),
        _block("p0004", 4, "Part I", "section"),
        _block("p0005", 5, "1. Name one river-valley civilization.", "major_question"),
        _block("p0006", 6, "a. Mention where people settled.", "subquestion", parent="p0005"),
        _block("p0007", 7, "b. Mention one crop they grew.", "subquestion", parent="p0005"),
        _block("p0008", 8, "2. Why did cities need walls?", "major_question"),
        _block("p0009", 9, "Part II", "section"),
        _block("p0010", 10, "Cities grew where water and farmland were reliable.", "passage", parent="p0008"),
        _block("p0011", 11, "3. Describe a marketplace.", "major_question"),
        _block("p0012", 12, "Part III", "section"),
        _block("p0013", 13, "4. What is a primary source?", "major_question"),
    ]
    assign_labels(blocks, policy)
    majors = [block for block in blocks if block.block_type == "major_question"]
    subs = [block for block in blocks if block.block_type == "subquestion"]
    assert [block.final_label for block in majors] == ["1.", "2.", "3.", "4."]
    assert [block.final_label for block in subs] == ["a.", "b."]
    assert all(not block.final_label.startswith("Q") for block in blocks)
    assert all(block.final_label not in {"A)", "B)", "a)", "b)"} for block in blocks)

    teacher = HISTORY
    once = tmp_path / "once.docx"
    format_exam(blocks, teacher, once, HISTORY)
    first = [paragraph.text for paragraph in Document(str(once)).paragraphs]
    assert any(text.startswith("1. Name one") for text in first)
    assert any(text.startswith("a. Mention where") for text in first)
    assert any(text.startswith("b. Mention one") for text in first)
    assert not any(text.startswith("Q") for text in first)
    assert not any(re.fullmatch(r"[A-D]\)", text.strip()) for text in first)
    assert not any("1. 1." in text or "a. a." in text for text in first)

    formatted = {block.source_id: block.original_text for block in blocks}
    for paragraph in first:
        if paragraph.startswith("1. Name"):
            formatted["p0005"] = paragraph
        elif paragraph.startswith("2. Why"):
            formatted["p0008"] = paragraph
        elif paragraph.startswith("3. Describe"):
            formatted["p0011"] = paragraph
        elif paragraph.startswith("4. What"):
            formatted["p0013"] = paragraph
        elif paragraph.startswith("a. Mention"):
            formatted["p0006"] = paragraph
        elif paragraph.startswith("b. Mention"):
            formatted["p0007"] = paragraph
    for block in blocks:
        block.original_text = formatted[block.source_id]
    assign_labels(blocks, policy)
    twice = tmp_path / "twice.docx"
    format_exam(blocks, teacher, twice, HISTORY)
    second = [paragraph.text for paragraph in Document(str(twice)).paragraphs]
    assert second == first
    assert not any("1. 1." in text or "a. a." in text for text in second)
